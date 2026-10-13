# -*- coding: utf-8 -*-
"""Markdown -> Word(.docx) 转换脚本。

流程：
  1. 复用 md2pdf.compress_images() 把 docs/images/ 降采样到 docs/images_compressed/
  2. 读取 docs/项目策划设计文档.md（跳过 YAML frontmatter，生成封面页）
  3. 逐行解析：标题 / 段落 / 表格 / 图片 / 图注 / 列表 / 引用 / 分割线
  4. 写出 docs/项目策划设计文档.docx，并写入文档属性（标题/作者/主题/关键词）

依赖：python-docx、Pillow（图片处理由 md2pdf 负责）。
"""
import re
import sys
from pathlib import Path

# 复用 md2pdf 的降采样与 frontmatter 解析
sys.path.insert(0, str(Path(__file__).resolve().parent))
import md2pdf as mp  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

MD_PATH = mp.ROOT / "docs" / "项目策划设计文档.md"
DST = mp.ROOT / "docs" / "项目策划设计文档.docx"
IMG_OUT = mp.IMG_OUT_DIR

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
CODE_FONT = "Consolas"
AZURE = RGBColor(0x8C, 0x1F, 0x28)   # 蜀锦红
GRAY = RGBColor(0x6B, 0x72, 0x80)


def set_font(run, name, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text, size=None):
    """处理行内的 **加粗**、*斜体*、`代码`。"""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            set_font(paragraph.add_run(part[2:-2]), BODY_FONT, size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            set_font(paragraph.add_run(part[1:-1]), CODE_FONT, size, color=AZURE)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            set_font(paragraph.add_run(part[1:-1]), BODY_FONT, size, italic=True)
        else:
            set_font(paragraph.add_run(part), BODY_FONT, size)
    return paragraph


def is_sep(s):
    return bool(re.match(r"^\|[\s:\-]+\|$", s.strip()))


def split_row(s):
    s = s.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def add_cover(doc, meta):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(str(meta.get("title", "") or "")), HEAD_FONT, 26, bold=True, color=AZURE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(str(meta.get("author", "") or "")), HEAD_FONT, 15)
    doc.add_page_break()


def build():
    # 0) 降采样图片
    st = mp.compress_images()
    mapping = st["mapping"]
    print("[0/2] 图片降采样：%d 张，%s -> %s" % (st["count"], mp.human(st["before"]), mp.human(st["after"])))

    text = MD_PATH.read_text(encoding="utf-8")
    meta, body = mp.split_frontmatter(text)
    lines = body.split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    for lvl in range(1, 5):
        key = "Heading %d" % lvl
        if key in [s.name for s in doc.styles]:
            hs = doc.styles[key]
            hs.font.name = HEAD_FONT
            hs.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
            hs.font.color.rgb = AZURE

    add_cover(doc, meta)

    tables = 0
    images = 0
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if s == "":
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(b) for b in block if not is_sep(b)]
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Table Grid"
                for ri, r in enumerate(rows):
                    cells = table.add_row().cells
                    for ci in range(cols):
                        txt = r[ci] if ci < len(r) else ""
                        p = cells[ci].paragraphs[0]
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.0
                        add_inline(p, txt, 10)
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
                tables += 1
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="Heading %d" % len(m.group(1)))
            add_inline(p, m.group(2))
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            name = Path(m.group(2)).name
            out = mapping.get(name, name)
            img_path = IMG_OUT / out
            if not img_path.exists():
                img_path = mp.IMG_DIR / name
            if img_path.exists():
                doc.add_picture(str(img_path), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                images += 1
            i += 1
            continue

        # 图注
        if re.match(r"^\*图 .*\*$", s):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            set_font(p.add_run(s.strip("*")), BODY_FONT, 9, italic=True, color=GRAY)
            i += 1
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            set_font(p.add_run(s.lstrip(">").strip()), BODY_FONT, 10, color=GRAY)
            i += 1
            continue

        # 分割线
        if re.match(r"^-{3,}$", s):
            i += 1
            continue

        # 无序列表
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:].strip())
            i += 1
            continue

        # 有序列表（保留原编号文本）
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            add_inline(p, s)
            i += 1
            continue

        # 普通段落
        add_inline(doc.add_paragraph(), s)
        i += 1

    cp = doc.core_properties
    cp.title = str(meta.get("title", "") or "")
    cp.author = str(meta.get("author", "") or "")
    cp.subject = mp.PDF_METADATA["subject"]
    cp.keywords = mp.PDF_METADATA["keywords"]

    doc.save(str(DST))
    print("[1/2] DOCX 已生成：%s" % DST)

    check = Document(str(DST))
    print("[2/2] 校验：paragraphs=%d tables=%d images=%d"
          % (len(check.paragraphs), len(check.tables), len(check.inline_shapes)))
    return 0


if __name__ == "__main__":
    sys.exit(build())
